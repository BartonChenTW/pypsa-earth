"""
Build ``pypsa_tw/data/taiwan_timeseries.csv``: Taiwan electricity data by year,
labelled as history, projection or target.

Inputs (``pypsa_tw/data/official/``):

- ``moeaea_generation_by_source_annual.csv``: Energy Administration, generation
  by source 2005-2025 (https://data.gov.tw/dataset/16481)
- ``moeaea_capacity_by_source_annual.csv``: capacity by source 2005-2025
  (https://data.gov.tw/dataset/16480)
- ``moeaea_energy_indicators_annual.csv``: energy indicators 2005-2025
  (https://data.gov.tw/dataset/8308)
- ``taipower_peak_load_by_year.csv``: Taipower peak load 1982-2025
  (https://data.gov.tw/dataset/8307)
- ``taiwan_projections_targets.csv``: projections and targets transcribed from the
  MOEA National Power Supply-Demand Report 113年度 (Tables 3-1, 3-2), plus policy targets

Every row carries a ``source_id`` (see ``sources.csv``: publisher, edition, links,
local copy, checksum) and a ``locator`` (table and page, or file column).
- ``moea_thermal_schedule_2024_2034.csv``: the same report's unit-by-unit plan of
  thermal additions and retirements (Figure 3-3)
- GEGIS demand projections shipped with PyPSA-Earth
  (``data/ssp2-2.6/<year>/era5_2013/Asia.csv``)
- ``moeaea_energy_balance_toe_1982_2025.xlsx``: Energy Administration energy balance
  (toe units, one sheet per year): energy use of every sector by fuel, transport by
  mode and fuel, industry by branch
- ``moeaea_co2_by_sector_1990_2025.csv``: CO2 from fuel combustion by sector, direct and
  with electricity allocated, transcribed from the two PDFs next to it
- ``moeaea_household_electricity_by_appliance_2024.csv``: household electricity by
  appliance (2024 survey, Residential Sector Energy Statistics 2024, Figures 12-13)

Run from the repository root:

    python pypsa_tw/data/build_timeseries.py
"""

from pathlib import Path

import pandas as pd

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
OFF = HERE / "official"
OUT = HERE / "taiwan_timeseries.csv"

# Source ids from sources.csv (the registry of every source used here, with its
# publisher, edition, links, local copy and checksum).
GEN, CAP, IND, PEAK = "moeaea_generation_annual", "moeaea_capacity_annual", "moeaea_indicators_annual", "taipower_peak_annual"
SOURCES_FILE = HERE / "sources.csv"

# (column stem in the Energy Administration files, series id, English, Chinese)
SOURCES = [
    ("火力_燃煤", "coal", "Coal", "燃煤"),
    ("火力_燃氣", "gas", "Gas", "燃氣"),
    ("火力_燃油", "oil", "Oil", "燃油"),
    ("核能", "nuclear", "Nuclear", "核能"),
    ("再生能源_慣常水力", "hydro", "Conventional hydro", "慣常水力"),
    ("再生能源_太陽光電", "solar", "Solar PV", "太陽光電"),
    ("再生能源_風力", "wind", "Wind", "風力"),
    ("再生能源_地熱", "geothermal", "Geothermal", "地熱"),
    ("再生能源_生質能", "biomass", "Biomass", "生質能"),
    ("再生能源_廢棄物", "waste", "Waste", "廢棄物"),
    ("抽蓄水力", "pumped_storage", "Pumped storage", "抽蓄水力"),
]

INDICATORS = [
    ("再生能源發電占總發電量比例(%)", "re_share", "Renewable share of generation", "再生能源發電占比", "%"),
    ("再生能源裝置容量占總裝置容量比例(%)", "re_capacity_share", "Renewable share of capacity", "再生能源裝置容量占比", "%"),
    ("電力排碳係數(公斤CO2e/度)", "grid_emission_factor", "Grid emission factor", "電力排碳係數", "kg CO2e/kWh"),
    ("燃料燃燒二氧化碳排放量(百萬公噸)", "co2_fuel_combustion", "CO2 from fuel combustion (all sectors)", "燃料燃燒二氧化碳排放量", "Mt"),
    ("負載率(%)", "load_factor", "System load factor", "負載率", "%"),
    ("線路損失率(%)", "line_losses", "Line losses", "線路損失率", "%"),
    ("進口能源依存度(%)", "import_dependence", "Imported energy dependence", "進口能源依存度", "%"),
]


# ---------- Energy use in all sectors ----------
# Energy balance (能源平衡表, toe units): one sheet per ROC year, rows = flows and
# sectors (item numbers), columns = products. 1 toe = 10^7 kcal = 11.63 MWh.
BAL, CO2S, HOME = "moeaea_energy_balance", "moeaea_co2_by_sector", "moeaea_household_appliances_2024"
BAL_FILE = OFF / "moeaea_energy_balance_toe_1982_2025.xlsx"
TOE_TWH = 11.63e-6
BAL_YEARS = range(2005, 2026)
# (series id, column position in a sheet row, column header in the sheet, English, Chinese)
BAL_FUELS = [
    ("coal", 3, "煤及煤產品", "Coal and coal products", "煤及煤產品"),
    ("oil", 15, "原油及石油產品", "Oil products", "石油產品"),
    ("gas", 39, "天然氣", "Natural gas", "天然氣"),
    ("biomass_waste", 42, "生質能及廢棄物", "Biomass and waste", "生質能及廢棄物"),
    ("electricity", 53, "電力", "Electricity", "電力"),
    ("solar_thermal", 54, "太陽熱能", "Solar thermal", "太陽熱能"),
    ("heat", 55, "熱能", "Heat (steam)", "熱能"),
]
BAL_TOTAL = (56, "總計")
BAL_PRODUCTS = {"gasoline": [(24, "車用汽油")], "diesel": [(30, "柴油")], "fuel_oil": [(31, "燃料油")],
                "jet": [(26, "航空汽油"), (27, "航空燃油-汽油型"), (28, "航空燃油-煤油型")]}
# (series id, item number, label check (prefix), English, Chinese). Residential and
# non-energy use are found by label: their item numbers changed in 2018 (94/95 -> 101/102).
BAL_SECTORS = [
    ("industry", 35, "工業部門", "Industry", "工業部門"),
    ("transport", 73, "運輸部門", "Transport", "運輸部門"),
    ("residential", None, "住宅部門", "Residential", "住宅部門"),
    ("services", 83, "服務業部門", "Services", "服務業部門"),
    ("agriculture", 80, "農業部門", "Agriculture", "農業部門"),
    ("non_energy", None, "非能源消費", "Non-energy use (feedstock)", "非能源消費（原料）"),
    ("energy_own_use", 22, "能源部門自用", "Energy sector own use", "能源部門自用"),
]
BAL_TRANSPORT_MODES = [("road", 75, "公路", "Road", "公路"), ("rail", 76, "鐵路", "Rail", "鐵路"),
                       ("domestic_air", 74, "國內航空", "Domestic aviation", "國內航空"),
                       ("domestic_water", 78, "國內水運", "Domestic shipping", "國內水運"),
                       ("pipeline", 77, "管線運輸", "Pipelines", "管線運輸")]
BAL_TRANSPORT_FUELS = [("gasoline", "Gasoline", "車用汽油"), ("diesel", "Diesel", "柴油"), ("jet", "Jet fuel", "航空燃油"),
                       ("fuel_oil", "Fuel oil", "燃料油"), ("electricity", "Electricity", "電力"),
                       ("other", "Other (LPG, gas, biofuels)", "其他（液化石油氣、天然氣、生質燃料）")]
BAL_INDUSTRY = [("chemicals", 43, "化學材料", "Chemicals and fertilisers", "化學材料及肥料"),
                ("basic_metals", 59, ("金屬基本工業", "基本金屬製造業"), "Basic metals (steel)", "基本金屬（鋼鐵）"),
                ("electronics", 65, ("電腦通信", "電子產品"), "Electronics and electrical equipment", "電子產品及電力設備"),
                ("non_metallic", 54, "非金屬礦物", "Cement, glass, ceramics", "非金屬礦物（水泥、玻璃、陶瓷）"),
                ("paper", 41, "紙漿", "Pulp and paper", "紙漿及紙製品"),
                ("textiles", 38, "紡織", "Textiles", "紡織成衣")]


def _label(v):
    """Row or column label; NFKC folds the CJK compatibility characters some sheets use (e.g. U+F9BE for 料)."""
    import unicodedata

    return unicodedata.normalize("NFKC", str(v)).strip() if v else ""


def _num(v):
    return 0.0 if v in (None, "") else float(v)


def read_energy_balance():
    """{year: {"rows": {item number: row}, "by_label": {label: row}}} for BAL_YEARS; checks the layout."""
    import openpyxl

    wb = openpyxl.load_workbook(BAL_FILE, read_only=True, data_only=True)
    years = {}
    for sheet in wb.sheetnames:
        year = int(sheet) + 1911
        if year not in BAL_YEARS:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        head = [_label(x) for x in rows[4]]
        for _, col, name, *_ in BAL_FUELS:
            assert head[col] == name, (sheet, col, head[col])
        assert head[BAL_TOTAL[0]] == BAL_TOTAL[1]
        for parts in BAL_PRODUCTS.values():
            for col, name in parts:
                assert head[col] == name, (sheet, col, head[col])
        by_no, by_label = {}, {}
        for r in rows[8:]:
            label = _label(r[2])
            if label.startswith("電能與熱能產出"):
                break
            if r[0] is not None and str(r[0]).strip().isdigit():
                by_no[int(str(r[0]).strip())] = r
            by_label.setdefault(label, r)
        years[year] = {"sheet": sheet, "rows": by_no, "by_label": by_label}
    return years


def _bal_row(y, no, check):
    r = y["rows"][no]
    label = _label(r[2])
    assert label.startswith(check if isinstance(check, str) else tuple(check)), (y["sheet"], no, label)
    return r


def energy_balance_rows():
    """Final energy use by sector and fuel, transport by mode and fuel, industry by branch (TWh)."""
    bal = read_energy_balance()
    rows = []
    scope = "Taiwan (national)"

    def add(series, en, zh, year, toe, locator, note=""):
        rows.append(row(series, en, zh, "history", year, toe * TOE_TWH, "TWh", scope, BAL, locator, note=note))

    sector_rows = {}
    for year, y in bal.items():
        sh = f"sheet {y['sheet']}"
        for sid, no, check, en, zh in BAL_SECTORS:
            r = _bal_row(y, no, check) if no else y["by_label"][check]
            sector_rows[(year, sid)] = r
            item = str(r[0]).strip()
            add(f"sector_energy_{sid}", f"Energy use: {en}", f"能源消費：{zh}", year, _num(r[BAL_TOTAL[0]]),
                f"{sh}, item {item} {check}, column 總計",
                note="Energy sector own use is not final consumption; shown for completeness" if sid == "energy_own_use" else
                     "Mainly naphtha and LPG used as petrochemical feedstock" if sid == "non_energy" else "")
        # Sectors add up to final consumption (item 33).
        final = _num(_bal_row(y, 33, "最終消費")[BAL_TOTAL[0]])
        parts = sum(_num(sector_rows[(year, s)][BAL_TOTAL[0]]) for s, *_ in BAL_SECTORS if s != "energy_own_use")
        assert abs(parts - final) <= 1e-6 * final + 1, (year, parts, final)

        # Fuel mix of each end-use sector.
        for sid, no, check, en, zh in BAL_SECTORS[:5]:
            r = sector_rows[(year, sid)]
            fuels = 0.0
            for fid, col, name, fen, fzh in BAL_FUELS:
                fuels += _num(r[col])
                add(f"sector_fuel_{sid}_{fid}", f"{en}: {fen}", f"{zh}：{fzh}", year, _num(r[col]),
                    f"{sh}, item {str(r[0]).strip()} {check}, column {name}")
            total = _num(r[BAL_TOTAL[0]])
            assert abs(fuels - total) <= 0.005 * total + 1, (year, sid, fuels, total)

        # Transport by mode and by fuel.
        tr = sector_rows[(year, "transport")]
        for mid, no, check, en, zh in BAL_TRANSPORT_MODES:
            add(f"transport_mode_{mid}", f"Transport: {en}", f"運輸：{zh}", year, _num(_bal_row(y, no, check)[BAL_TOTAL[0]]),
                f"{sh}, item {no} {check}, column 總計")
        named = 0.0
        for fid, en, zh in BAL_TRANSPORT_FUELS:
            if fid == "other":
                v, where = _num(tr[BAL_TOTAL[0]]) - named, "column 總計 minus the fuels above"
            elif fid == "electricity":
                v, where = _num(tr[53]), "column 電力"
            else:
                v = sum(_num(tr[c]) for c, _ in BAL_PRODUCTS[fid])
                where = "column " + " + ".join(n for _, n in BAL_PRODUCTS[fid])
            named += v if fid != "other" else 0
            add(f"transport_fuel_{fid}", f"Transport fuel: {en}", f"運輸燃料：{zh}", year, v, f"{sh}, item 73 運輸部門, {where}")

        # Industry by branch.
        ind = _num(sector_rows[(year, "industry")][BAL_TOTAL[0]])
        named = 0.0
        for bid, no, check, en, zh in BAL_INDUSTRY:
            r = _bal_row(y, no, check)
            named += _num(r[BAL_TOTAL[0]])
            add(f"industry_branch_{bid}", f"Industry: {en}", f"工業：{zh}", year, _num(r[BAL_TOTAL[0]]),
                f"{sh}, item {no} {_label(r[2])}, column 總計")
        add("industry_branch_other", "Industry: Other branches", "工業：其他行業", year, ind - named,
            f"{sh}, item 35 工業部門 minus the branches above")

    out = pd.DataFrame(rows)
    # Drop series that are zero in every year (e.g. solar thermal in some sectors).
    keep = out.groupby("series").value.transform(lambda v: v.abs().max() > 1e-6)
    return out[keep].to_dict(orient="records")


def co2_sector_rows():
    """CO2 from fuel combustion by sector, direct and with electricity allocated to its users (Mt)."""
    co2 = pd.read_csv(OFF / "moeaea_co2_by_sector_1990_2025.csv")
    names = {"energy": ("Energy sector", "能源部門"), "industry": ("Industry", "工業部門"), "transport": ("Transport", "運輸部門"),
             "agriculture": ("Agriculture", "農業部門"), "services": ("Services", "服務業部門"),
             "residential": ("Residential", "住宅部門"), "total": ("Total", "合計")}
    method = {"direct": ("direct", "直接排放", "Direct emissions: power plants count in the energy sector"),
              "incl_indirect": ("incl. electricity", "含電力間接排放", "Electricity emissions allocated to the sectors that use the power")}
    rows = []
    for _, r in co2.iterrows():
        men, mzh, note = method[r.method]
        for sid, (en, zh) in names.items():
            rows.append(row(f"co2_sector_{r.method}_{sid}", f"CO2 by sector ({men}): {en}", f"部門別 CO2（{mzh}）：{zh}",
                            "history", int(r.year), r[sid] / 1e3, "Mt", "Taiwan (national)", CO2S,
                            f"{r.table}, year {int(r.year)}, column {zh}; transcribed in official/moeaea_co2_by_sector_1990_2025.csv",
                            note=note))
    return rows


def household_appliance_rows():
    """Household electricity use by appliance, 2024 survey (whole year and summer months)."""
    a = pd.read_csv(OFF / "moeaea_household_electricity_by_appliance_2024.csv")
    assert abs(a.whole_year_pct.sum() - 100) < 0.05 and abs(a.summer_pct.sum() - 100) < 0.05
    rows = []
    for _, r in a.iterrows():
        note = r.note if isinstance(r.note, str) else ""
        rows.append(row(f"home_elec_year_{r.appliance}", f"Household electricity, whole year: {r.appliance_en}",
                        f"家庭全年用電：{r.appliance_zh}", "history", 2024, r.whole_year_pct, "%", "Households (survey)", HOME,
                        "Figure 12, pp. 24-25 (PDF p. 14)", note=note))
        rows.append(row(f"home_elec_summer_{r.appliance}", f"Household electricity, summer (Jun-Sep): {r.appliance_en}",
                        f"家庭夏月用電：{r.appliance_zh}", "history", 2024, r.summer_pct, "%", "Households (survey)", HOME,
                        "Figure 13, pp. 26-27 (PDF p. 15)", note=""))
    return rows


def load_sources():
    """Source registry; checks that each local copy still matches its recorded checksum."""
    import hashlib

    src = pd.read_csv(SOURCES_FILE, dtype=str).fillna("").set_index("source_id")
    for sid, r in src.iterrows():
        if r.local_file and r.sha256:
            digest = hashlib.sha256((HERE / r.local_file).read_bytes()).hexdigest()
            assert digest == r.sha256, f"{sid}: {r.local_file} changed since it was registered"
    return src


REG = load_sources()


def row(series, en, zh, kind, year, value, unit, scope, source_id, locator="", note="", evidence=None):
    src = REG.loc[source_id]
    return {"series": series, "indicator_en": en, "indicator_zh": zh, "kind": kind, "year": year,
            "value": round(float(value), 4), "unit": unit, "scope": scope, "source_id": source_id,
            "source_name": src.short_cite, "locator": locator, "link": src.landing_url,
            "evidence": evidence or src.evidence, "note": note}


def main():
    rows = []
    gen = pd.read_csv(OFF / "moeaea_generation_by_source_annual.csv", encoding="utf-8-sig").set_index("西元年")
    cap = pd.read_csv(OFF / "moeaea_capacity_by_source_annual.csv", encoding="utf-8-sig").set_index("西元年")
    ind = pd.read_csv(OFF / "moeaea_energy_indicators_annual.csv", encoding="utf-8-sig").set_index("西元年")
    peak = pd.read_csv(OFF / "taipower_peak_load_by_year.csv", encoding="utf-8-sig").set_index("年度")

    for year in gen.index:
        rows.append(row("generation_total", "Total generation", "總發電量", "history", year,
                        gen.loc[year, "總發電量_全國(統計數值)"] / 1e3, "TWh", "Taiwan (national)", GEN, "column 總發電量_全國(統計數值)",
                        note="Gross generation including self-generation"))
        tpc = gen.loc[year, "總發電量_台電(統計數值)"] + gen.loc[year, "總發電量_民營電廠(統計數值)"]
        rows.append(row("generation_taipower_ipp", "Generation, Taipower + IPPs", "台電及民營電廠發電量", "history", year,
                        tpc / 1e3, "TWh", "Taipower system", GEN, "columns 總發電量_台電 + 總發電量_民營電廠 (統計數值)",
                        note="Taipower-owned plus independent power producers; excludes power bought from self-generators"))
        for stem, sid, en, zh in SOURCES:
            col = f"{stem}_全國(統計數值)"
            if col in gen.columns:
                rows.append(row(f"generation_{sid}", f"Generation: {en}", f"發電量：{zh}", "history", year,
                                gen.loc[year, col] / 1e3, "TWh", "Taiwan (national)", GEN, f"column {col}"))
    for year in cap.index:
        rows.append(row("capacity_total", "Installed capacity", "裝置容量", "history", year,
                        cap.loc[year, "總發電裝置容量_全國(統計數值)"] / 1e3, "GW", "Taiwan (national)", CAP, "column 總發電裝置容量_全國(統計數值)",
                        note="Includes self-generation"))
        for stem, sid, en, zh in SOURCES:
            col = f"{stem}_全國(統計數值)"
            if col in cap.columns:
                rows.append(row(f"capacity_{sid}", f"Capacity: {en}", f"裝置容量：{zh}", "history", year,
                                cap.loc[year, col] / 1e3, "GW", "Taiwan (national)", CAP, f"column {col}"))
    for year in ind.index:
        for col, sid, en, zh, unit in INDICATORS:
            if pd.notna(ind.loc[year, col]):
                rows.append(row(sid, en, zh, "history", year, ind.loc[year, col], unit, "Taiwan (national)", IND, f"column {col}"))
    for year in peak.index:
        rows.append(row("peak_load", "Peak load (day)", "尖峰負載（日間）", "history", year,
                        peak.loc[year, "尖峰負載(MW)"] / 1e3, "GW", "Taipower system", PEAK, "column 尖峰負載(MW)"))
        rows.append(row("reserve_margin", "Reserve margin", "備用容量率", "history", year,
                        peak.loc[year, "備用容量率(％)"], "%", "Taipower system", PEAK, "column 備用容量率(％)",
                        note="Night-time reserve margin from 2022" if year >= 2022 else ""))

    # Projections and targets transcribed from the report (numeric years only here).
    pt = pd.read_csv(OFF / "taiwan_projections_targets.csv", dtype={"year": str})
    for _, r in pt.iterrows():
        if not r.year.isdigit():
            continue
        try:
            value = float(r.value)
        except ValueError:
            continue
        unit, v = r.unit, value
        if unit == "MW":
            unit, v = "GW", value / 1e3
        series = {"re_capacity_solar": "capacity_solar", "re_capacity_offwind": "capacity_offshore_wind",
                  "re_capacity_onwind": "capacity_onshore_wind", "re_capacity_geothermal": "capacity_geothermal",
                  "re_capacity_biomass": "capacity_biomass_waste", "re_capacity_hydro": "capacity_hydro",
                  "re_capacity_total": "capacity_renewables"}.get(r.series, r.series)
        rows.append(row(series, r.indicator_en, r.indicator_zh, r.kind, int(r.year), v, unit, r.scope,
                        r.source_id, r.locator if isinstance(r.locator, str) else "",
                        note=r.note if isinstance(r.note, str) else "", evidence=r.evidence))

    # Thermal plan by year (Figure 3-3 of the report, transcribed unit by unit).
    sched = pd.read_csv(OFF / "moea_thermal_schedule_2024_2034.csv")
    fuel = {"CCGT": ("gas", "gas", "燃氣"), "Hard Coal": ("coal", "coal", "燃煤"), "Oil": ("oil", "oil", "燃油")}
    for (action, ft, year), g in sched.groupby(["action", "fueltype", "year"]):
        sid, en, zh = fuel[ft]
        verb_en, verb_zh = ("additions", "新增") if action == "add" else ("retirements", "除役")
        rows.append(row(f"planned_{action}_{sid}", f"Planned {en} {verb_en}", f"{zh}機組{verb_zh}規劃", "projection",
                        int(year), g.mw.sum() / 1e3, "GW", "Taipower system", "moea_psd_fy2024",
                        "Figure 3-3, p. 18 (PDF p. 22); transcribed in official/moea_thermal_schedule_2024_2034.csv",
                        note="Units: " + ", ".join(f"{u} ({m:g} MW, month {mo})"
                                                   for u, m, mo in zip(g.unit_en, g.mw, g.month))))

    # PyPSA-Earth default demand projection (GEGIS, SSP2-2.6, 2013 weather).
    for year in [2030, 2040, 2050]:
        f = REPO / "data" / "ssp2-2.6" / str(year) / "era5_2013" / "Asia.csv"
        if not f.exists():
            continue
        d = pd.read_csv(f, sep=";")
        tw = d[d.region_code.astype(str) == "TW"]["Electricity demand"]
        rows.append(row("gegis_demand", "Electricity demand (PyPSA-Earth default)", "電力需求（PyPSA-Earth 預設）",
                        "projection", year, tw.sum() / 1e6, "TWh", "Taiwan (national)", "gegis_ssp2_26",
                        "data/ssp2-2.6/<year>/era5_2013/Asia.csv, region TW, sum of hourly values",
                        note="Unscaled GEGIS projection; the model scales the 2030 profile by 0.749"))
        rows.append(row("gegis_peak", "Peak demand (PyPSA-Earth default)", "尖峰需求（PyPSA-Earth 預設）",
                        "projection", year, tw.max() / 1e3, "GW", "Taiwan (national)", "gegis_ssp2_26",
                        "data/ssp2-2.6/<year>/era5_2013/Asia.csv, region TW, maximum hourly value",
                        note="Hourly peak of the unscaled GEGIS profile"))

    # Official demand path: the report's 1.7%/yr growth applied to 2024 national generation.
    base = gen.loc[2024, "總發電量_全國(統計數值)"] / 1e3
    for year in range(2025, 2035):
        rows.append(row("generation_forecast", "Generation, official growth rate applied", "發電量（依官方成長率推估）",
                        "projection", year, base * 1.017 ** (year - 2024), "TWh", "Taiwan (national)",
                        "derived_moea_growth", "1.7%/yr: moea_psd_fy2024, Section 3.1(1), p. 7 (PDF p. 11); 2024 base: moeaea_generation_annual",
                        note="Not published as such: the report's average growth rate applied to 2024 national generation"))

    # Energy use in all sectors, CO2 by sector, household end uses.
    rows += energy_balance_rows() + co2_sector_rows() + household_appliance_rows()

    out = pd.DataFrame(rows).sort_values(["series", "year", "kind"])
    out.to_csv(OUT, index=False, encoding="utf-8")
    print(f"wrote {len(out)} rows, {out.series.nunique()} series to {OUT.relative_to(REPO)}")
    print(out.groupby("kind").size().to_string())


if __name__ == "__main__":
    main()
